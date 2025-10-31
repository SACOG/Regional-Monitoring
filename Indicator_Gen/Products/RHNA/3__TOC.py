



import openpyxl
from openpyxl.styles import Font, Border, Side, PatternFill, Alignment, Protection
from pathlib import Path


path_sp   = Path.home() / 'Sacramento Area Council of Governments\Regional Monitoring and Reporting - Documents'
path_prod = path_sp    / 'Products' / 'RHNA'
path_raw  = path_prod  / 'New Data Collected'

path_git = Path.home() / 'Documents' / 'Projects' / 'Regional-Monitoring' / 'Indicator_Gen'
path_census = path_git  / 'Data' / 'Census'
path_config0 = path_git  / 'config'

path_prod = path_git / 'Products' / 'RHNA'
path_config = path_prod / 'config'
path_yaml = path_config / 'RHNA_indicators.yaml'
path_func = path_config / 'RHNA_functions.py'
path_py = path_prod / 'python'
path_i = Path(r'I:\Projects\Josh\RHNA')

# path_out = Path(r'I:\Projects\Josh\RHNA\Final Products')
path_out = Path(r'C:\Users\jfontes\Documents\Projects\General\RHNA\Final Products')

path_geo = Path(r'I:\Projects\Josh\Geospatial Data\crosswalks')
path_lodes = Path(r'I:\Projects\Josh\Regional Monitoring')





counties = ['El Dorado', 'Placer', 'Sacramento', 'Sutter', 'Yolo', 'Yuba']

dt_juris = {
    'El Dorado': ['Placerville', 'South Lake Tahoe', 'Unincorporated'],
    'Placer': ['Auburn', 'Colfax', 'Lincoln', 'Loomis', 'Rocklin', 'Roseville', 'Unincorporated'],
    'Sacramento': ['Citrus Heights', 'Elk Grove', 'Folsom', 'Galt', 'Isleton', 'Rancho Cordova', 'Sacramento', 'Unincorporated'],
    'Sutter': ['Live Oak', 'Yuba City', 'Unincorporated'],
    'Yolo': ['Davis', 'West Sacramento', 'Winters', 'Woodland', 'Unincorporated'],
    'Yuba': ['Marysville', 'Wheatland', 'Unincorporated']
}


print(); print()

for county in counties:
    print(county)
    for jurisdiction in dt_juris[county]:

        # Define the file paths
        source_workbook_path = path_config / 'TOC.xlsx'
        destination_workbook_path = path_out / county.replace(' County', '') / jurisdiction / f'RHNA_{jurisdiction}.xlsx'


        # Load the source workbook and get the sheet to copy
        source_wb = openpyxl.load_workbook(source_workbook_path)
        source_sheet = source_wb.active  # or source_wb['YourSheetName'] if you know the sheet name

        # Load the destination workbook
        destination_wb = openpyxl.load_workbook(destination_workbook_path)

        # Get the 'Sheet' from the destination workbook
        if 'Sheet' in destination_wb.sheetnames: # Sheet
            destination_sheet = destination_wb['Sheet']
        else:
            destination_sheet = destination_wb.active

        # Clear the destination sheet
        for row in destination_sheet.iter_rows():
            for cell in row:
                cell.value = None


        # Copy values and styles
        for row in source_sheet.iter_rows():
            for cell in row:
                new_cell = destination_sheet.cell(row=cell.row, column=cell.column, value=cell.value)

                # Copy font
                new_cell.font = Font(
                    name=cell.font.name,
                    size=cell.font.size,
                    bold=cell.font.bold,
                    italic=cell.font.italic,
                    vertAlign=cell.font.vertAlign,
                    underline=cell.font.underline,
                    strike=cell.font.strike,
                    color=cell.font.color
                )

                # Copy border
                new_cell.border = Border(
                    left=cell.border.left,
                    right=cell.border.right,
                    top=cell.border.top,
                    bottom=cell.border.bottom,
                    diagonal=cell.border.diagonal,
                    diagonal_direction=cell.border.diagonal_direction,
                    outline=cell.border.outline,
                    vertical=cell.border.vertical,
                    horizontal=cell.border.horizontal
                )

                # Copy fill
                new_cell.fill = PatternFill(
                    fill_type=cell.fill.fill_type,
                    start_color=cell.fill.start_color,
                    end_color=cell.fill.end_color
                )

                # Copy alignment
                new_cell.alignment = Alignment(
                    horizontal=cell.alignment.horizontal,
                    vertical=cell.alignment.vertical,
                    text_rotation=cell.alignment.text_rotation,
                    wrap_text=cell.alignment.wrap_text,
                    shrink_to_fit=cell.alignment.shrink_to_fit,
                    indent=cell.alignment.indent
                )

                # Copy protection
                new_cell.protection = Protection(
                    locked=cell.protection.locked,
                    hidden=cell.protection.hidden
                )


        # Copy merged cells
        for merged_range in source_sheet.merged_cells.ranges:
            destination_sheet.merge_cells(str(merged_range))


        destination_sheet.column_dimensions['A'].width = 12
        destination_sheet.column_dimensions['B'].width = 75
        destination_sheet.column_dimensions['C'].width = 52
        destination_sheet.column_dimensions['D'].width = 25
        destination_sheet.column_dimensions['E'].width = 65
        destination_sheet.column_dimensions['F'].width = 25
        destination_sheet.column_dimensions['G'].width = 150


        for row in range(1, destination_sheet.max_row + 1):
            destination_sheet.row_dimensions[row].height = 16


        # Rename the sheet
        destination_sheet.title = 'Indicators List'

        # Save the destination workbook
        destination_wb.save(destination_workbook_path)

