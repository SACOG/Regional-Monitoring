


import pandas as pd
from pathlib import Path
from tqdm import tqdm
import traceback
from docx import Document

PATH_CONFIG = Path.home() / 'Documents' / 'Projects' / 'Regional-Monitoring' / 'Indicator_Gen' / 'Products' / 'RHNA' / 'config'

import sys
sys.path.append(str(PATH_CONFIG))
import rhna
FILE_YAML = rhna.load_yaml()

PATH_PROD = Path.home() / 'Documents' / 'Projects' / 'General' / 'RHNA' / 'Final Products'



def doc_add_table(df):
    table = document.add_table(rows=1, cols=df.shape[1])
    table.style = 'Table Grid'

    hdr_cells = table.rows[0].cells
    for i, col_name in enumerate(df.columns):
        hdr_cells[i].text = str(col_name)

    for index, row in df.iterrows():
        row_cells = table.add_row().cells
        for i, cell_value in enumerate(row):
            row_cells[i].text = str(cell_value)





rhna.print3()

if __name__ == '__main__':

    indicators = ['HSG_2']
    
    for folder in PATH_PROD.iterdir():
        rhna.print2()
        county = folder.stem; print(county); print()
        path_county = PATH_PROD / county
        for folder in path_county.iterdir():
            jurisdiction = folder.stem; print(jurisdiction)
            path_juris = path_county / jurisdiction

            file_doc    = path_juris / f'RHNA_{jurisdiction}.docx'
            file_wkbk   = path_juris / f'RHNA_{jurisdiction}.xlsx'
            path_plots  = path_juris / 'plots'
            path_tables = path_juris / 'tables'


            document = Document()
            document.add_heading(f'{jurisdiction}, {county} County', level=0)


            for indicator in tqdm(indicators):
                try:
    
                    theme  = FILE_YAML[indicator]['Theme' ][0]
                    title  = FILE_YAML[indicator]['Title' ][0]
                    source = FILE_YAML[indicator]['Source'][0]

                    df_table = pd.read_excel(file_wkbk, sheet_name=indicator, skiprows=2)
                    df_table = df_table.dropna()
                    
                    # file_tables = [file for file in path_tables.glob('*.csv') if f'{indicator}_' in str(file)]

                    # if len(file_tables) == 1:
                    #     df_prod = pd.read_csv(file_tables[0])
                    # else:
                    #     df_pct  = pd.read_csv(file_tables[0])
                    #     df_prod = pd.read_csv(file_tables[1])
                    #     df_pct  = df_pct .melt(id_vars = 'Geography', var_name = 'Variable', value_name='Percentage')
                    #     df_prod = df_prod.melt(id_vars = 'Geography', var_name = 'Variable', value_name='Total'     )
                    #     df_prod = df_prod.merge(df_pct)
                    #     # df_prod = df_prod.pivot_table(index='Variable', columns='Geography', values=['Total', 'Percentage']).reset_index()
                    #     breakpoint()

                    if '1' in indicator:
                        document.add_heading(f'{theme}', level=1)
                    document.add_paragraph('This is the first paragraph of text.')
                    document.add_paragraph('Here is some more text for the second paragraph.')

                    p = document.add_paragraph()
                    p.add_run(f'{indicator}: {title}').bold = True


                    doc_add_table(df_prod)

                    p = document.add_paragraph(f'Source: {source}')
                    document.add_page_break()


                except Exception as e: print(e); traceback.print_exc(); print()

            document.save(file_doc)
