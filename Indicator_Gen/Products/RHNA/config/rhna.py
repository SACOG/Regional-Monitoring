

def print2(): print();print()
def print3(): print();print();print()




EXPORT=True



import numpy as np
import pandas as pd
from pathlib import Path
import os
import requests
import gzip
import io
from tqdm import tqdm
from datetime import datetime
import time
import re
import yaml
import traceback
import plotly.express as px
from IPython.display import display

import openpyxl
from openpyxl.drawing.image import Image
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.styles import Font
# from openpyxl.styles import numbers
from openpyxl.styles import Border, Side
border_thin = Side(style='thin')
# pip install kaleido==0.1.0.post1


PATH_DATA    = Path(r'I:\Projects\Josh\RHNA\Data')
PATH_GIT     = Path(__file__).parent.parent.parent.parent
PATH_CENSUS  = PATH_GIT / 'Data' / 'Census'
PATH_CONFIG0 = PATH_GIT / 'config'
PATH_PROD    = PATH_GIT / 'Products' / 'RHNA'
PATH_CONFIG  = PATH_GIT / 'Products' / 'RHNA' / 'config'
PATH_PY      = PATH_GIT / 'Products' / 'RHNA' / 'python'
PATH_OUT = Path.home() / 'Documents' / 'Projects' / 'Local' / 'RHNA' / 'Final Products'
PATH_GEO = Path(r'I:\Projects\Josh\Geospatial Data')
PATH_LODES = Path(r'I:\Projects\Josh\Regional Monitoring')


def load_yaml():
    try:
        with open(PATH_CONFIG / 'rhna.yaml', 'r', encoding='utf-8') as yaml_file:
            yaml_rhna = yaml.load(yaml_file, Loader=yaml.SafeLoader)
    except FileNotFoundError:
        print(f"Error: The file at {PATH_CONFIG} does not exist.")
    
    return yaml_rhna


def split_notes(indicator):
    FILE_YAML = load_yaml()
    notes = FILE_YAML[indicator]['Notes']
    
    lines = notes.split('\\n')
    rows_new = [{'Indicator': 'Notes' if i == 0 else '', indicator: line} for i, line in enumerate(lines) if line]
    
    df_notes = pd.DataFrame(rows_new)
    
    return df_notes



def import_gz_from_url(url):

    response = requests.get(url, stream=True)
    response.raise_for_status()  # Raise an exception for non-200 status codes
    compressed_file = io.BytesIO(response.content)
    decompressed_file = gzip.GzipFile(fileobj=compressed_file)
    data = decompressed_file.read() # Read the decompressed data

    return data




def sort_categorical(df, col, sort_list):
    df['Sort'] = pd.Categorical(df[col], sort_list)
    df = df.sort_values('Sort').drop('Sort', axis=1)
    return df



## DOF ---

def dof_import(indicator):
    
    workbooks = os.listdir(PATH_DATA)
    workbooks = [workbook for workbook in workbooks if indicator in workbook]
    print(); print()
    print('Workbooks to import: ', workbooks)     

    path_places   = PATH_DATA / f'RHNA_{indicator} Places DOF.xlsx'  
    path_counties = PATH_DATA / f'RHNA_{indicator} Counties DOF.xlsx'
    path_mpo      = PATH_DATA / f'RHNA_{indicator} MPO DOF.xlsx'     
    
    df_places   = pd.read_excel(path_places  , sheet_name='Places'  )
    df_counties = pd.read_excel(path_counties, sheet_name='Counties')
    df_mpo      = pd.read_excel(path_mpo     , sheet_name='MPO'     )

    return df_places, df_counties, df_mpo


def dof_clean(df_places, df_counties, df_mpo):

    df_places   = df_places  [df_places  ['MPO'] == 'SACOG']
    df_counties = df_counties[df_counties['MPO'] == 'SACOG']
    df_mpo      = df_mpo[     df_mpo     ['MPO'] == 'SACOG']
    df_mpo['MPO'] = df_mpo['MPO'] + ' Region'

    df_places   = df_places  [[          'County', 'Jurisdiction', 'Year', 'Population']]
    df_counties = df_counties[[          'County',                 'Year', 'Population']]
    df_mpo      = df_mpo     [['MPO'   ,                           'Year', 'Population']]

    df_places   = df_places  .rename(columns={'Jurisdiction':'Geography'})
    df_counties = df_counties.rename(columns={'County'      :'Geography'})
    df_mpo      = df_mpo     .rename(columns={'MPO'         :'Geography'})

    df_places   = df_places  .reset_index(drop=True)
    df_counties = df_counties.reset_index(drop=True)
    df_mpo      = df_mpo     .reset_index(drop=True)

    return df_places, df_counties, df_mpo


def dof_sub(df_places, df_counties, county):
    
    df_counties_sub = df_counties[df_counties['Geography'] == county]
    df_places_sub   = df_places  [df_places  ['County'   ] == county]
        
    df_places_sub   = df_places_sub  .reset_index(drop=True)
    df_counties_sub = df_counties_sub.reset_index(drop=True)

    return df_places_sub, df_counties_sub

def dof_index(df_places_sub, df_counties_sub, df_mpo, jurisdiction):

    df_plot1 = df_places_sub.copy()
    df_plot2 = df_counties_sub.copy()
    df_plot3 = df_mpo.copy()

    df_plot1 = df_plot1[df_plot1['Geography'] == jurisdiction]
    df_plot1 = df_plot1[df_plot1['Population'] > 0]
    base_year = df_plot1['Year'].min()
    df_year = df_plot1[df_plot1['Year'] == base_year]
    df_year = df_year.rename(columns={'Population':'base_year'})
    df_year = df_year.drop('Year', axis=1)
    df_plot1 = df_plot1.merge(df_year, on=['County', 'Geography'], how='left')
    df_plot1['growth'] = df_plot1['Population']/df_plot1['base_year']-1
    df_plot1 = df_plot1.drop('base_year', axis=1)

    df_plot2 = df_plot2[df_plot2['Year'] >= base_year]
    df_year = df_plot2[df_plot2['Year'] == base_year]
    df_year = df_year.rename(columns={'Population':'base_year'})
    df_year = df_year.drop('Year', axis=1)
    df_plot2 = df_plot2.merge(df_year, on=['Geography'], how='left')
    df_plot2['growth'] = df_plot2['Population']/df_plot2['base_year']-1
    df_plot2 = df_plot2.drop('base_year', axis=1)

    df_plot3 = df_plot3[df_plot3['Year'] >= base_year]
    df_year = df_plot3[df_plot3['Year'] == base_year]
    df_year = df_year.rename(columns={'Population':'base_year'})
    df_year = df_year.drop('Year', axis=1)
    df_plot3 = df_plot3.merge(df_year, on=['Geography'], how='left')
    df_plot3['growth'] = df_plot3['Population']/df_plot3['base_year']-1
    df_plot3 = df_plot3.drop(['base_year'], axis=1)

    return df_plot1, df_plot2, df_plot3



## CHAS ---


def chas_import(PATH_DATA, indicator):
    
    workbooks = os.listdir(PATH_DATA)
    workbooks = [workbook for workbook in workbooks if f'RHNA_{indicator}' in workbook]
    print(); print()
    print('Workbooks to import: ', workbooks)
    
    path_places   = PATH_DATA / f'RHNA_{indicator} Places ACS5.xlsx'  
    path_counties = PATH_DATA / f'RHNA_{indicator} Counties ACS5.xlsx'
    path_mpo      = PATH_DATA / f'RHNA_{indicator} MPO ACS5.xlsx'     
    
    df_places   = pd.read_excel(path_places  , sheet_name='Places'  )
    df_counties = pd.read_excel(path_counties, sheet_name='Counties')
    df_mpo      = pd.read_excel(path_mpo     , sheet_name='MPO'     )

    return df_places, df_counties, df_mpo



def chas_clean(df_places, df_counties, df_mpo, columns, values):
    
    df_places  ['NAME'] = df_places  ['NAME'].str.replace(' CDP, California' , '', regex=True)
    df_places  ['NAME'] = df_places  ['NAME'].str.replace(' city, California', '', regex=True)
    df_counties['NAME'] = df_counties['NAME'].str.replace(', California'     , '', regex=True)
    df_mpo['MPO'] = df_mpo['MPO'] + ' Region'
    
    df_places   = df_places  .rename(columns={'NAME':'Geography'})
    df_counties = df_counties.rename(columns={'NAME':'Geography'})
    df_mpo      = df_mpo     .rename(columns={'MPO' :'Geography'})
    df_mpo = df_mpo.drop_duplicates()

    df_places   = df_places  [df_places  ['Year'] == df_places  ['Year'].max()]
    df_counties = df_counties[df_counties['Year'] == df_counties['Year'].max()]
    df_mpo      = df_mpo     [df_mpo     ['Year'] == df_mpo     ['Year'].max()]
    
    df_places   = df_places  [['County Name', 'Geography', columns, values, 'Percent']]
    df_counties = df_counties[[               'Geography', columns, values, 'Percent']]
    df_mpo      = df_mpo     [[               'Geography', columns, values, 'Percent']]

    df_places   = df_places  .reset_index(drop=True)
    df_counties = df_counties.reset_index(drop=True)
    df_mpo      = df_mpo     .reset_index(drop=True)

    return df_places, df_counties, df_mpo


def chas_sub(df_places, df_counties, county):
    
    df_counties_sub = df_counties[df_counties['Geography'  ] == county                       ]
    df_places_sub   = df_places  [df_places  ['County Name'] == county.replace(' County', '')]

    df_places_sub = df_places_sub.drop('County Name', axis=1)
    df_places_sub   = df_places_sub  .reset_index(drop=True)
    df_counties_sub = df_counties_sub.reset_index(drop=True)

    return df_places_sub, df_counties_sub


def chas_pivot(indicator, df_places_sub, county, jurisdiction, columns, values, df_counties_sub=None, df_mpo=None):

    if indicator in ['POPEMP_10', 'POPEMP_19', 'POPEMP_20', 'POPEMP_21', 'POPEMP_22', 'POPEMP_25', 'HSG_1', 'HSG_5', 'HSG_6'
                        , 'OVER_4', 'OVER_5', 'OVER_6', 'OVER_8', 'OVER_9', 'FARM_2', 'LGFEM_1', 'LGFEM_3', 'LGFEM_4', 'LGFEM_5'
                        , 'SEN_1', 'SEN_2', 'SEN_3', 'DISAB_3', 'HOMELS_1', 'HOMELS_2', 'HOMELS_3', 'HOMELS_4', 'ELI_2', 'AFFH_1', 'AFFH_2']:

        df_plot = df_places_sub[df_places_sub['Geography'] == jurisdiction]
        df_prod = df_plot.pivot_table(index=['Geography', 'Variable'], columns=columns, values=values).reset_index()

        if indicator == 'POPEMP_10':
            vars_to_sort = ['75k or more', '50k to 75k', ' 25k to 50k', '10k to 25k', 'Less than 10k']
        if indicator == 'POPEMP_19':
            vars_to_sort = ['Moved in 2021 or later', 'Moved in 2018 to 2020', 'Moved in 2010 to 2017', 'Moved in 2000 to 2009', 'Moved in 1999 or earlier']

        df_prod['Sort'] = pd.Categorical(df_prod['Variable'], vars_to_sort)
        df_prod = df_prod.sort_values(['Sort'], ascending=[False])
        df_prod = df_prod.drop(['Geography', 'Sort'], axis=1)
        df_pct = df_plot.pivot_table(index=['Geography', 'Variable'], columns=columns, values='Percent').reset_index()
        df_pct['Sort'] = pd.Categorical(df_prod['Variable'], vars_to_sort)
        df_pct = df_pct.sort_values(['Sort'], ascending=[False])
        df_pct = df_pct.drop(['Geography', 'Sort'], axis=1)
    
    elif indicator in ['HSG_4', 'HSG_11', 'OVER_1', 'OVER_3', 'SEN_4', 'DISAB_1', 'DISAB_4', 'DISAB_5', 'ELI_3']:
        pass
    elif indicator in ['POPEMP_19']:
        pass
        
    else:
        df_prod = pd.concat([df_places_sub[df_places_sub['Geography'] == jurisdiction], df_counties_sub, df_mpo])
        if 'Percent' in df_places_sub.columns:
            df_prod = df_prod.drop('Percent', axis=1)
        df_prod = df_prod.pivot_table(index='Geography', columns=columns, values=values).reset_index()
        df_prod['Sort'] = pd.Categorical(df_prod['Geography'], [jurisdiction, county, 'SACOG Region'])
        df_prod = df_prod.sort_values(['Sort'])
        df_prod = df_prod.drop(['Sort'], axis=1)

        if 'Percent' in df_places_sub.columns:
            df_pct = pd.concat([df_places_sub[df_places_sub['Geography'] == jurisdiction], df_counties_sub, df_mpo])
            df_pct = df_pct.drop(values, axis=1)
            
            df_pct = df_pct.pivot_table(index='Geography', columns=columns, values='Percent').reset_index()
            df_pct['Sort'] = pd.Categorical(df_pct['Geography'], [jurisdiction, county, 'SACOG Region'])
            df_pct = df_pct.sort_values(['Sort'])
            df_pct = df_pct.drop(['Sort'], axis=1)
        
    if jurisdiction == 'Sacramento':
        print()
        print('Final Product:')
        display(df_prod.head())
        print()

    return df_prod, df_pct
        

## HDMA ---

def hmda_import(years, counties, dtypes):
    '''
    Years can only be from 2018, to 2024. County codes can be found by using the online tool: https://ffiec.cfpb.gov/data-browser/data/2023?category=counties. 
    '''
    print(''); print('Importing data from HDMA...'); print('')

    list_df = []

    for year in tqdm(years):
        if int(year) < 2018 or int(year) > int(datetime.now().year - 1): 
            print(f"Error: the year {year} does not have accessible data.") 
            continue

        url = 'https://ffiec.cfpb.gov/v2/data-browser-api/view/csv?counties=' + ','.join(counties) + '&years=' + str(year)
        
        df = pd.read_csv(url, dtype=dtypes)

        list_df.append(df)
    
    df = pd.concat(list_df, ignore_index=True)

    return df


## CHP ---

def chp_sub(df_places, df_counties, county):
    
    df_counties_sub = df_counties[df_counties['Geography'  ] == county]
    df_places_sub   = df_places  [df_places  ['County'     ] == county]

    df_places_sub = df_places_sub.drop('County', axis=1)
    df_places_sub   = df_places_sub  .reset_index(drop=True)
    df_counties_sub = df_counties_sub.reset_index(drop=True)

    return df_places_sub, df_counties_sub


def chp_clean(df_chp):

    df_chp['MPO'] = 'SACOG Region'
    df_chp.loc[df_chp['Risk Level'] == 'HIgh', 'Risk Level'] = 'High'
    df_chp['County'] = df_chp['County'] + ' County'

    df_mpo      = df_chp.groupby(['MPO'           , 'Risk Level'], as_index=False).agg(affordable_units=('Affordable Units', 'sum')).rename(columns={'MPO':'Geography'})
    df_counties = df_chp.groupby(['County'        , 'Risk Level'], as_index=False).agg(affordable_units=('Affordable Units', 'sum')).rename(columns={'County':'Geography'})
    df_places   = df_chp.groupby(['County', 'City', 'Risk Level'], as_index=False).agg(affordable_units=('Affordable Units', 'sum')).rename(columns={'City':'Geography'})

    df_mpo2      = df_chp.groupby(['MPO'           ], as_index=False).agg(affordable_units=('Affordable Units', 'sum')).rename(columns={'MPO':'Geography'})
    df_counties2 = df_chp.groupby(['County'        ], as_index=False).agg(affordable_units=('Affordable Units', 'sum')).rename(columns={'County':'Geography'})
    df_places2   = df_chp.groupby(['County', 'City'], as_index=False).agg(affordable_units=('Affordable Units', 'sum')).rename(columns={'City':'Geography'})

    df_mpo2     ['Risk Level'] = 'Total Assisted Units in Database'
    df_counties2['Risk Level'] = 'Total Assisted Units in Database'
    df_places2  ['Risk Level'] = 'Total Assisted Units in Database'

    df_mpo      = pd.concat([df_mpo     , df_mpo2     ])
    df_counties = pd.concat([df_counties, df_counties2])
    df_places   = pd.concat([df_places  , df_places2  ])

    df_mpo      = df_mpo     .merge(df_mpo2     .drop('Risk Level', axis=1).rename(columns={'affordable_units':'total_affordable_units'}), on=[          'Geography'], how='left')
    df_counties = df_counties.merge(df_counties2.drop('Risk Level', axis=1).rename(columns={'affordable_units':'total_affordable_units'}), on=[          'Geography'], how='left')
    df_places   = df_places  .merge(df_places2  .drop('Risk Level', axis=1).rename(columns={'affordable_units':'total_affordable_units'}), on=['County', 'Geography'], how='left')

    df_mpo     ['Percent'] = df_mpo     ['affordable_units']/df_mpo     ['total_affordable_units']
    df_counties['Percent'] = df_counties['affordable_units']/df_counties['total_affordable_units']
    df_places  ['Percent'] = df_places  ['affordable_units']/df_places  ['total_affordable_units']

    df_mpo      = df_mpo     .drop('total_affordable_units', axis=1)
    df_counties = df_counties.drop('total_affordable_units', axis=1)
    df_places   = df_places  .drop('total_affordable_units', axis=1)

    df_mpo     ['Sort'] = pd.Categorical(df_mpo     ['Risk Level'], ['Low', 'Moderate', 'High', 'Very High', 'Total Assisted Units in Database'])
    df_counties['Sort'] = pd.Categorical(df_counties['Risk Level'], ['Low', 'Moderate', 'High', 'Very High', 'Total Assisted Units in Database'])
    df_places  ['Sort'] = pd.Categorical(df_places  ['Risk Level'], ['Low', 'Moderate', 'High', 'Very High', 'Total Assisted Units in Database'])

    df_mpo      = df_mpo     .sort_values(['Geography', 'Sort']).reset_index(drop=True).drop('Sort', axis=1)
    df_counties = df_counties.sort_values(['Geography', 'Sort']).reset_index(drop=True).drop('Sort', axis=1)
    df_places   = df_places  .sort_values(['Geography', 'Sort']).reset_index(drop=True).drop('Sort', axis=1)

    return df_places, df_counties, df_mpo



def chp_pivot(df_places_sub, county, jurisdiction, columns, values, df_counties_sub=None, df_mpo=None):

    df_prod = pd.concat([df_places_sub[df_places_sub['Geography'] == jurisdiction], df_counties_sub, df_mpo])
    if 'Percent' in df_places_sub.columns:
        df_prod = df_prod.drop('Percent', axis=1)
    df_prod = df_prod.pivot_table(index='Geography', columns=columns, values=values).reset_index()
    df_prod['Sort'] = pd.Categorical(df_prod['Geography'], [jurisdiction, county, 'SACOG Region'])
    df_prod = df_prod.sort_values(['Sort']).drop(['Sort'], axis=1)

    if 'Percent' in df_places_sub.columns:
        df_pct = pd.concat([df_places_sub[df_places_sub['Geography'] == jurisdiction], df_counties_sub, df_mpo])
        df_pct = df_pct.drop(values, axis=1)
        
        df_pct = df_pct.pivot_table(index='Geography', columns=columns, values='Percent').reset_index()
        df_pct['Sort'] = pd.Categorical(df_pct['Geography'], [jurisdiction, county, 'SACOG Region'])
        df_pct = df_pct.sort_values(['Sort']).drop(['Sort'], axis=1)

    if jurisdiction == 'Sacramento':
        print()
        print('Final Product:')
        display(df_prod.head())
        print()

    return df_prod, df_pct






## ACS ---------------------------------------------------------------------------------------------------------------------------------------------




def cols_geo_rename(df, cols, county, jurisdiction):
    if 'County' not in county:
        county = f'{county} County'
    cols = [col.replace('[county]', county) for col in cols]
    cols = [col.replace('[jurisdiction]', jurisdiction) for col in cols]
    df = df[cols]
    return df


def acs_import(indicator):
    
    workbooks = os.listdir(PATH_DATA)
    workbooks = [workbook for workbook in workbooks if indicator in workbook]
    print(); print()
    print('Workbooks to import: ', workbooks)

    if indicator == 'ELI_4':
        path_places   = PATH_DATA / f'RHNA_{indicator} Places DP5.xlsx'
        path_counties = PATH_DATA / f'RHNA_{indicator} Counties DP5.xlsx'
        path_mpo      = PATH_DATA / f'RHNA_{indicator} MPO DP5.xlsx'
    else:
        path_places   = PATH_DATA / f'RHNA_{indicator} Places ACS5.xlsx'  
        path_counties = PATH_DATA / f'RHNA_{indicator} Counties ACS5.xlsx'
        path_mpo      = PATH_DATA / f'RHNA_{indicator} MPO ACS5.xlsx'     
    
    df_places   = pd.read_excel(path_places  , sheet_name='Places'  )
    df_counties = pd.read_excel(path_counties, sheet_name='Counties')
    df_mpo      = pd.read_excel(path_mpo     , sheet_name='MPO'     )

    df_places   = df_places  .rename(columns={'Race_Ethnicity':'Race/Ethnicity', 'Percentage':'Percent'})
    df_counties = df_counties.rename(columns={'Race_Ethnicity':'Race/Ethnicity', 'Percentage':'Percent'})
    df_mpo      = df_mpo     .rename(columns={'Race_Ethnicity':'Race/Ethnicity', 'Percentage':'Percent'})

    if indicator == 'OVER_3':
        df_places   = df_places  [df_places  ['Variable']=='More than occupant per room'] # More than one occupant per room
        df_counties = df_counties[df_counties['Variable']=='More than occupant per room'] # More than one occupant per room
        df_mpo      = df_mpo     [df_mpo     ['Variable']=='More than occupant per room'] # More than one occupant per room
        

    if len(df_places['Race/Ethnicity'].unique()) > 1:
        df_places   = df_places  [df_places  ['Race/Ethnicity'] != 'All']
        df_counties = df_counties[df_counties['Race/Ethnicity'] != 'All']
        df_mpo      = df_mpo     [df_mpo     ['Race/Ethnicity'] != 'All']

    return df_places, df_counties, df_mpo



def acs_clean(df_places, df_counties, df_mpo, params):

    params = params['prod']
    
    df_places  ['NAME'] = df_places  ['NAME'].str.replace(' CDP, California' , '', regex=True)
    df_places  ['NAME'] = df_places  ['NAME'].str.replace(' city, California', '', regex=True)
    df_counties['NAME'] = df_counties['NAME'].str.replace(', California'     , '', regex=True)
    df_mpo['MPO'] = df_mpo['MPO'] + ' Region'
    
    df_places   = df_places  .rename(columns={'NAME':'Geography'})
    df_counties = df_counties.rename(columns={'NAME':'Geography'})
    df_mpo      = df_mpo     .rename(columns={'MPO' :'Geography'})
    df_mpo = df_mpo.drop_duplicates()

    if 'Variable' in df_places.columns and params['var'] != 'Race/Ethnicity':
        df_places   = df_places  .rename(columns={'Variable':params['var']})
        df_counties = df_counties.rename(columns={'Variable':params['var']})
        df_mpo      = df_mpo     .rename(columns={'Variable':params['var']})
    
    df_places   = df_places  [df_places  ['Year'] == df_places  ['Year'].max()][['County Name', 'Geography', params['var'], params['value'], 'Percent']].reset_index(drop=True)
    df_counties = df_counties[df_counties['Year'] == df_counties['Year'].max()][[               'Geography', params['var'], params['value'], 'Percent']].reset_index(drop=True)
    df_mpo      = df_mpo     [df_mpo     ['Year'] == df_mpo     ['Year'].max()][[               'Geography', params['var'], params['value'], 'Percent']].reset_index(drop=True)

    if 'Race/Ethnicity' in df_places.columns:
        df_places  .loc[df_places  ['Race/Ethnicity'].isin(['Some other race (NH)', 'Two or more races (NH)']), 'Race/Ethnicity'] = 'Other race or multiple races (NH)'
        df_counties.loc[df_counties['Race/Ethnicity'].isin(['Some other race (NH)', 'Two or more races (NH)']), 'Race/Ethnicity'] = 'Other race or multiple races (NH)'
        df_mpo     .loc[df_mpo     ['Race/Ethnicity'].isin(['Some other race (NH)', 'Two or more races (NH)']), 'Race/Ethnicity'] = 'Other race or multiple races (NH)'


    return df_places, df_counties, df_mpo


def acs_sub(df_places, df_counties, county):
    
    df_counties_sub = df_counties[df_counties['Geography'  ] == county].reset_index(drop=True)
    df_places_sub   = df_places  [df_places  ['County Name'] == county.replace(' County', '')].drop('County Name', axis=1).reset_index(drop=True)

    return df_places_sub, df_counties_sub


def acs_pivot(indicator, df_places_sub, county, jurisdiction, columns, values, variable=None, df_counties_sub=None, df_mpo=None):

    if indicator in ['POPEMP_10', 'POPEMP_18', 'POPEMP_19', 'POPEMP_20', 'POPEMP_21', 'POPEMP_22', 'POPEMP_25', 'HSG_1', 'HSG_5', 'HSG_6'
                        , 'OVER_4', 'OVER_5', 'OVER_6', 'OVER_8', 'OVER_9', 'FARM_2', 'LGFEM_1', 'LGFEM_3', 'LGFEM_4', 'LGFEM_5'
                        , 'SEN_1', 'SEN_2', 'SEN_3', 'DISAB_3', 'HOMELS_1', 'HOMELS_2', 'HOMELS_3', 'HOMELS_4', 'ELI_2', 'ELI_3', 'AFFH_1', 'AFFH_2']:

        df_plot = df_places_sub[df_places_sub['Geography'] == jurisdiction]
        df_prod = df_plot.pivot_table(index=['Geography', variable], columns=columns, values=values).reset_index()

        if indicator == 'POPEMP_10': vars_to_sort = ['Less than $10k', '$10k to $25k', '$25k to $50k', '$50k to $75k', '$75k or more']
        if indicator == 'POPEMP_18': vars_to_sort = ['Age 15-24', 'Age 25-34', 'Age 35-44', 'Age 45-54', 'Age 55-59', 'Age 60-64', 'Age 65-74', 'Age 75-84', 'Age 85+']
        if indicator == 'POPEMP_19': vars_to_sort = ['Moved in 1999 or earlier', 'Moved in 2000 to 2009', 'Moved in 2010 to 2017', 'Moved in 2018 to 2020', 'Moved in 2021 or later']
        if indicator == 'POPEMP_22': vars_to_sort = ['Detached single-family homes', 'Attached single-family homes', 'Multi-family housing', 'Mobile homes', 'Boat, RV, van, or other']
        if indicator == 'HSG_5': vars_to_sort = ['0 bedrooms', '1 bedroom', '2 bedrooms', '3-4 bedrooms', '5 or more bedrooms']
        if indicator == 'HSG_6': vars_to_sort = ['Lacking kitchen facilities', 'Lacking plumbing facilities']
        if indicator == 'OVER_6': vars_to_sort = ['0%-30% of income used for housing', '30%-50% of income used for housing', '50% or more of income used for housing', 'Not computed']
        if indicator == 'LGFEM_1': vars_to_sort = ['1 person household', '2 person household', '3 person household', '4 person household', '5 or more person household']
        if indicator == 'LGFEM_4': vars_to_sort = ['Married-couple family', 'Female-headed family household', 'Male-headed family household', 'Householders living alone', 'Other non-family households']
        if indicator == 'LGFEM_5': vars_to_sort = ['Female-headed households without children', 'Female-headed households with children']
        if indicator == 'SEN_2': vars_to_sort = ['Age 0-17', 'Age 18-64', 'Age 65+']
        if indicator == 'DISAB_3': vars_to_sort = ['With a disability', 'No disability']
        if indicator in ['POPEMP_20', 'ELI_3']: vars_to_sort = ['American Indian or Alaska Native', 'Asian', 'Black or African American', 'Native Hawaiian or other Pacific Islander', 'Hispanic or Latino', 'Some other race', 'Two or more races', 'White (NH)']

        df_prod['Sort'] = pd.Categorical(df_prod[variable], vars_to_sort)
        df_prod = df_prod.sort_values(['Sort']).drop(['Geography', 'Sort'], axis=1)

        df_pct = df_plot.pivot_table(index=['Geography', variable], columns=columns, values='Percent').reset_index()
        df_pct['Sort'] = pd.Categorical(df_pct[variable], vars_to_sort)
        df_pct = df_pct.sort_values(['Sort'])
        df_pct = df_pct.drop(['Geography', 'Sort'], axis=1)
    
    elif indicator in ['HSG_4', 'HSG_11', 'OVER_1', 'SEN_4', 'DISAB_1', 'DISAB_4', 'DISAB_5', 'ELI_3']:
        pass
    elif indicator in ['POPEMP_19']:
        pass
        
    else:

        df_prod = pd.concat([df_places_sub[df_places_sub['Geography'] == jurisdiction], df_counties_sub, df_mpo])
        if 'Percent' in df_places_sub.columns:
            df_prod = df_prod.drop('Percent', axis=1)
        df_prod = df_prod.pivot_table(index='Geography', columns=columns, values=values).reset_index()
        df_prod['Sort'] = pd.Categorical(df_prod['Geography'], [jurisdiction, county, 'SACOG Region'])
        df_prod = df_prod.sort_values(['Sort'])
        df_prod = df_prod.drop(['Sort'], axis=1)
        if indicator == 'HSG_7':
            df_prod = df_prod[['Geography', 'Units valued less than 250k', 'Units valued 250k-500k', 'Units valued 500k-750k', 'Units valued 750k-1M', 'Units valued 1M-1.5M', 'Units valued 1.5M-2M', 'Units valued 2M+']]
            df_prod = df_prod.rename(columns = {  'Units valued less than 250k': 'Units valued less than $250k'
                                                , 'Units valued 250k-500k': 'Units valued $250k-$500k'
                                                , 'Units valued 500k-750k': 'Units valued $500k-$750k'
                                                , 'Units valued 750k-1M': 'Units valued $750k-$1M'
                                                , 'Units valued 1M-1.5M': 'Units valued $1M-$1.5M'
                                                , 'Units valued 1.5M-2M': 'Units valued $1.5M-$2M'
                                                , 'Units valued 2M+': 'Units valued $2M+'
                                                })
        if indicator == 'HSG_9':
            df_prod = df_prod[['Geography', 'Rent less than 500', 'Rent 500-1,000', 'Rent 1,000-1,500', 'Rent 1,500-2,000', 'Rent 2,000-2,500', 'Rent 2,500-3,000', 'Rent 3,000 or more']]
            df_prod = df_prod.rename(columns = {'Rent less than 500': 'Rent less than $500'
                                            , 'Rent 500-1,000': 'Rent $500-$1,000'
                                            , 'Rent 1,000-1,500': 'Rent $1,000-$1,500'
                                            , 'Rent 1,500-2,000': 'Rent $1,500-$2,000'
                                            , 'Rent 2,000-2,500': 'Rent $2,000-$2,500'
                                            , 'Rent 2,500-3,000': 'Rent $2,500-$3,000'
                                            , 'Rent 3,000 or more': 'Rent $3,000 or more'
                                            })

        if 'Percent' in df_places_sub.columns:
            df_pct = pd.concat([df_places_sub[df_places_sub['Geography'] == jurisdiction], df_counties_sub, df_mpo])
            df_pct = df_pct.drop(values, axis=1)
            
            df_pct = df_pct.pivot_table(index='Geography', columns=columns, values='Percent').reset_index()
            df_pct['Sort'] = pd.Categorical(df_pct['Geography'], [jurisdiction, county, 'SACOG Region'])
            df_pct = df_pct.sort_values(['Sort'])
            df_pct = df_pct.drop(['Sort'], axis=1)
            if indicator == 'HSG_7':
                df_pct = df_pct[['Geography', 'Units valued less than 250k', 'Units valued 250k-500k', 'Units valued 500k-750k', 'Units valued 750k-1M', 'Units valued 1M-1.5M', 'Units valued 1.5M-2M', 'Units valued 2M+']]
                df_pct = df_pct.rename(columns = {'Units valued less than 250k': 'Units valued less than $250k'
                                                , 'Units valued 250k-500k': 'Units valued $250k-$500k'
                                                , 'Units valued 500k-750k': 'Units valued $500k-$750k'
                                                , 'Units valued 750k-1M': 'Units valued $750k-$1M'
                                                , 'Units valued 1M-1.5M': 'Units valued $1M-$1.5M'
                                                , 'Units valued 1.5M-2M': 'Units valued $1.5M-$2M'
                                                , 'Units valued 2M+': 'Units valued $2M+'
                                                })
            if indicator == 'HSG_9':
                df_pct = df_pct[['Geography', 'Rent less than 500', 'Rent 500-1,000', 'Rent 1,000-1,500', 'Rent 1,500-2,000', 'Rent 2,000-2,500', 'Rent 2,500-3,000', 'Rent 3,000 or more']]
                df_pct = df_pct.rename(columns = {'Rent less than 500': 'Rent less than $500'
                                                , 'Rent 500-1,000': 'Rent $500-$1,000'
                                                , 'Rent 1,000-1,500': 'Rent $1,000-$1,500'
                                                , 'Rent 1,500-2,000': 'Rent $1,500-$2,000'
                                                , 'Rent 2,000-2,500': 'Rent $2,000-$2,500'
                                                , 'Rent 2,500-3,000': 'Rent $2,500-$3,000'
                                                , 'Rent 3,000 or more': 'Rent $3,000 or more'
                                                })
        
    if jurisdiction == 'Sacramento':
        print()
        print('Final Product:')
        display(df_prod.head())
        print()

    return df_prod, df_pct




def main_acs(indicator, params):

    df_places, df_counties, df_mpo = acs_import(indicator)
    df_places, df_counties, df_mpo = acs_clean(df_places, df_counties, df_mpo, params)

    if indicator in ['POPEMP_3']:
        df_places   = df_places  .groupby(['County Name', 'Geography', 'Race/Ethnicity'], as_index=False)['Population'].sum()
        df_counties = df_counties.groupby([               'Geography', 'Race/Ethnicity'], as_index=False)['Population'].sum()
        df_mpo      = df_mpo     .groupby([               'Geography', 'Race/Ethnicity'], as_index=False)['Population'].sum()
        df_places  ['Percent'] = df_places  ['Population'] / df_places  .groupby(['County Name', 'Geography'])['Population'].transform('sum')
        df_counties['Percent'] = df_counties['Population'] / df_counties.groupby([               'Geography'])['Population'].transform('sum')
        df_mpo     ['Percent'] = df_mpo     ['Population'] / df_mpo     .groupby([               'Geography'])['Population'].transform('sum')

    if indicator in ['OVER_7']:
        df_places  .loc[df_places  ['Cost Burden']=='50% or more of income used for housing', 'Cost Burden'] = '50%+ of income used for housing'
        df_counties.loc[df_counties['Cost Burden']=='50% or more of income used for housing', 'Cost Burden'] = '50%+ of income used for housing'
        df_mpo     .loc[df_mpo     ['Cost Burden']=='50% or more of income used for housing', 'Cost Burden'] = '50%+ of income used for housing'

    counties = df_counties['Geography'].unique()

    for county in counties:

        print2()
        print(county)
        time.sleep(2)

        df_places_sub, df_counties_sub = acs_sub(df_places, df_counties, county)
        jurisdictions = df_places_sub['Geography'].unique()
        
        for jurisdiction in tqdm(jurisdictions, position=0):
            
            tqdm.write(jurisdiction)

            df_prod = pd.concat([df_places_sub[df_places_sub['Geography'] == jurisdiction], df_counties_sub, df_mpo]).pivot_table(index=params['prod']['index'], columns=params['prod']['pivot'], values=params['prod']['value']).reset_index()
            df_pct  = pd.concat([df_places_sub[df_places_sub['Geography'] == jurisdiction], df_counties_sub, df_mpo]).pivot_table(index=params['prod']['index'], columns=params['prod']['pivot'], values='Percent'              ).reset_index()
            
            df_prod = cols_geo_rename(df_prod, params['Columns'].split('<>'), county, jurisdiction)
            df_pct  = cols_geo_rename(df_pct , params['Columns'].split('<>'), county, jurisdiction)
            
            if indicator in ['HSG_7']:
                sort_list = ['Units valued less than 250k', 'Units valued 250k-500k', 'Units valued 500k-750k', 'Units valued 750k-1M', 'Units valued 1M-1.5M', 'Units valued 1.5M-2M', 'Units valued 2M+']
                df_prod = sort_categorical(df_prod, 'Home Value', sort_list)
                df_pct  = sort_categorical(df_pct , 'Home Value', sort_list)
            if indicator in ['HSG_9']:
                sort_list = ["Rent less than 500", "Rent 500-1,000", "Rent 1,000-1,500", "Rent 1,500-2,000", "Rent 2,000-2,500", "Rent 2,500-3,000", "Rent 3,000 or more"]
                df_prod = sort_categorical(df_prod, 'Rent Price', sort_list)
                df_pct  = sort_categorical(df_pct , 'Rent Price', sort_list)

            df_plot = pd.concat([df_places_sub[df_places_sub['Geography'] == jurisdiction], df_counties_sub, df_mpo])
            df_plot[f'Percent of {params['prod']['value']}'] = round(df_plot['Percent']*100, 1)
            df_plot['Sort'] = pd.Categorical(df_plot['Geography'], [jurisdiction, county, 'SACOG Region'])
            df_plot = df_plot.sort_values(['Sort', f'Percent of {params['prod']['value']}'], ascending=[True, False]).drop('Sort', axis=1)

            fig = make_fig(indicator, params, df_plot, county, jurisdiction)
            plot_rhna(fig, county, jurisdiction, indicator, params)
            export_rhna(county, jurisdiction, indicator, params, df_prod, df_pct)






## Plotting/Exporting --------------------------------------------------------------------------------------------------------------


def sort_plot(df_plot, indicator, jurisdiction):

    if indicator in ['POPEMP_21', 'POPEMP_23', 'POPEMP_24', 'HSG_2', 'HSG_7', 'HSG_9', 'HSG_11', 'OVER_3', 'OVER_4', 'OVER_5', 'OVER_7',
                     'OVER_8', 'LGFEM_1', 'LGFEM_2', 'LGFEM_3', 'LGFEM_4', 'LGFEM_5', 'SEN_1', 'SEN_3', 'SEN_4', 'DISAB_1', 'DISAB_5',
                     'HOMELS_1', 'HOMELS_2', 'ELI_2', 'ELI_3', 'ELI_4', 'AFFH_1', 'AFFH_3']:
        if indicator == 'POPEMP_23':
            df_plot['Household Type'] = df_plot['Household Type'].str.replace(' households', '')
        if indicator == 'POPEMP_24':
            var_map = {'Households with no children': 'No children'
                       , 'Households with 1 or more children under 18': "One or more children under 18"}
            df_plot['Household Type'] = df_plot['Household Type'].map(var_map)
        if indicator == 'HSG_2':
            var_map = {'Occupied housing units': 'Occupied'
                       , 'Vacant housing units': 'Vacant'}
            df_plot['Occupancy Status'] = df_plot['Occupancy Status'].map(var_map)
        if indicator == 'HSG_7':
            df_plot['Sort'] = pd.Categorical(df_plot['Home Value'], ['Units valued less than 250k', 'Units valued 250k-500k', 'Units valued 500k-750k', 'Units valued 750k-1M', 'Units valued 1M-1.5M', 'Units valued 1.5M-2M', 'Units valued 2M+'])
            df_plot = df_plot.sort_values('Sort').drop('Sort', axis=1)
            var_map = {'Units valued less than 250k': '&#36;0k-&#36;250k'
                        , 'Units valued 250k-500k': '&#36;250k-&#36;500k'
                        , 'Units valued 500k-750k': '&#36;500k-&#36;750k'
                        , 'Units valued 750k-1M': '&#36;750k-&#36;1M'
                        , 'Units valued 1M-1.5M': '&#36;1M-&#36;1.5M'
                        , 'Units valued 1.5M-2M': '&#36;1.5M-&#36;2M'
                        , 'Units valued 2M+': '&#36;2M+'}
            df_plot['Home Value'] = df_plot['Home Value'].map(var_map)
            df_plot['Share of Owner-Occupied Units'] = df_plot['Percent of Households'].copy()
        if indicator == 'HSG_9':
            df_plot['Sort'] = pd.Categorical(df_plot['Rent Price'], ["Rent less than 500", "Rent 500-1,000", "Rent 1,000-1,500", "Rent 1,500-2,000", "Rent 2,000-2,500", "Rent 2,500-3,000", "Rent 3,000 or more"])
            df_plot = df_plot.sort_values('Sort').drop('Sort', axis=1)
            var_map = {"Rent less than 500": "&#36;0-&#36;500"
                        , "Rent 500-1,000": "&#36;500-&#36;1,000"
                        , "Rent 1,000-1,500": "&#36;1,000-&#36;1,500"
                        , "Rent 1,500-2,000": "&#36;1,500-&#36;2,000"
                        , "Rent 2,000-2,500": "&#36;2,000-&#36;2,500"
                        , "Rent 2,500-3,000": "&#36;2,500-&#36;3,000"
                        , "Rent 3,000 or more": "&#36;3,000+"}
            df_plot['Rent Price'] = df_plot['Rent Price'].map(var_map)
            df_plot['Share of Renter-Occupied Units'] = df_plot['Percent of Households'].copy()
        if indicator == 'HSG_11':
            df_plot['Sort'] = pd.Categorical(df_plot['Income Group'], ["Very Low", "Low", "Moderate", "Above Moderate"])
            df_plot = df_plot.sort_values('Sort').drop('Sort', axis=1)
        if indicator in ['OVER_3', 'ELI_3']:
            var_map = {
                'American Indian or Alaska Native':'American Indian or<br>Alaska Native'
                , 'Asian':'Asian'
                , 'Black or African American':'Black or<br>African American'
                , 'Hispanic or Latino':'Hispanic or<br>Latino'
                , 'Native Hawaiian or other Pacific Islander':'Native Hawaiian or<br>other Pacific Islander'
                , 'Other race or multiple races':'Other race or<br>multiple races'
                , 'White (NH)': 'White (NH)'
            }
            df_plot['Race/Ethnicity'] = df_plot['Race/Ethnicity'].map(var_map)
            df_plot = df_plot[df_plot['Geography']==jurisdiction]
        if indicator in ['POPEMP_21', 'OVER_4', 'OVER_5', 'LGFEM_3', 'SEN_1', 'SEN_3']:
            var_map = {
                '0%-30% of AMI':'Less than 30%'
                , '31%-50% of AMI':'31%-50%'
                , '51%-80% of AMI':'51%-80%'
                , '81%-100% of AMI':'81%-100%'
                , 'Greater than 100% of AMI':'Greater than 100%'
            }
            df_plot['Income Level (% of AMI)'] = df_plot['Income Level'].map(var_map)
        if indicator == 'LGFEM_1':
            var_map = {
                '1 person household': '1 Person<br>Household',
                '2 person household': '2 Person<br>Household',
                '3 person household': '3 Person<br>Household',
                '4 person household': '4 Person<br>Household',
                '5 or more person household': '5 or More Person<br>Household'
            }
            df_plot['Household Size'] = df_plot['Household Size'].map(var_map)
        if indicator == 'LGFEM_2':
            df_plot['Sort'] = pd.Categorical(df_plot['Household Size'], ["1 person households", "2 person households", "3-4 person Households", "5 or more person households"])
            df_plot = df_plot.sort_values('Sort').drop('Sort', axis=1)
        if indicator == 'LGFEM_4':
            var_map = {
                'Married-couple family': 'Married-Couple<br>Family',
                'Female-headed family household': 'Female-Headed<br>Family',
                'Male-headed family household': 'Male-Headed<br>Family',
                'Householders living alone': 'Householders<br>Living Alone',
                'Other non-family households': 'Other Non-Family'
            }
            df_plot['Household Type'] = df_plot['Household Type'].map(var_map)
        if indicator == 'LGFEM_5':
            var_map = {
                'Female-headed households without children': 'Without Children',
                'Female-headed households with children': 'With Children'
            }
            df_plot['Family Status'] = df_plot['Family Status'].map(var_map)
        if indicator in ['SEN_4', 'DISAB_1']:
            var_map = {
                'With an ambulatory difficulty':'Ambulatory<br>Difficulty',
                'With an independent living difficulty':'Independent Living<br>Difficulty',
                'With a hearing difficulty':'Hearing<br>Difficulty',
                'With a self-care difficulty':'Self-Care<br>Difficulty',
                'With a cognitive difficulty':'Cognitive<br>Difficulty',
                'With a vision difficulty':'Vision<br>Difficulty'
            }
            df_plot['Disability Type'] = df_plot['Disability Type'].map(var_map)
        if indicator == 'DISAB_5':
            var_map = {
                'Home of Parent /Family /Guardian':'Home of<br>Parent/Family/Guardian'
                , 'Community Care Facility':'Community<br>Care<br>Facility'
                , 'Independent /Supported Living':'Independent/Supported<br>Living'
                , 'Foster /Family Home':'Foster/Family<br>Home'
                , 'Intermediate Care Facility':'Intermediate Care<br>Facility'
                , 'Other': 'Other'
            }
            df_plot['Residence Type'] = df_plot['Residence Type'].map(var_map)
        if indicator == 'HOMELS_1':
            var_map = {
                    'Persons in households with at least one adult and one child':'At Least One Adult<br>and One Child',
                    'Persons in households with only children':'Only Children',
                    'Persons in households without children':'Without Children',
                }
            df_plot['Household Type'] = df_plot['Household Type'].map(var_map)
        if indicator == 'HOMELS_2':
            var_map = {
                    'American Indian, Alaska Native, or Indigenous':'American Indian,<br>Alaska Native,<br>or Indigenous'
                    , 'Asian or Asian American':'Asian or<br>Asian American'
                    , 'Black, African American, or African':'Black,<br>African American,<br>or African'
                    , 'Hispanic/Latina/e/o Only':'Hispanic/Latina/e/o Only'
                    , 'Native Hawaiian or Other Pacific Islander':'Native Hawaiian or<br>Other Pacific Islander'
                    , 'Other race or multiple races':'Other Race or<br>Multiple Races'
                    , 'White': 'White'
                }
            df_plot['Race/Ethnicity'] = df_plot['Race/Ethnicity'].map(var_map)
            var_map = {
                    'Homeless Population (%)':'Share of homeless population'
                    , 'Overall Population (%)':'Share of overall population'
                }
            df_plot['Prop'] = df_plot['Prop'].map(var_map)
        if indicator == 'HOMELS_3':
            var_map = {
                    'Sheltered - Emergency Shelter':'Emergency Shelter'
                    , 'Sheltered - Transitional Housing':'Transitional Housing'
                    , 'Unsheltered':'Unsheltered'
                }
            df_plot['Shelter Status'] = df_plot['Shelter Status'].map(var_map)
        if indicator in ['OVER_8', 'ELI_2', 'AFFH_1']:
            var_map = {
                'American Indian or Alaska Native (NH)':'American Indian or<br>Alaska Native (NH)'
                , 'Asian (NH)':'Asian (NH)'
                , 'Black or African American (NH)':'Black or<br>African American (NH)'
                , 'Hispanic or Latino':'Hispanic or<br>Latino'
                , 'Native Hawaiian or other Pacific Islander (NH)':'Native Hawaiian or<br>other Pacific Islander (NH)'
                , 'Other race or multiple races (NH)':'Other race or<br>multiple races (NH)'
                , 'White (NH)': 'White (NH)'
            }
            df_plot['Race/Ethnicity'] = df_plot['Race/Ethnicity'].map(var_map)
        if indicator == 'ELI_4':
            var_map = {
                'Acutely low income': 'Acutely Low'
                , 'Extremely low income': 'Extremely Low'
                , 'Very low income': 'Very Low'
                , 'Lower income': 'Lower'
                , 'Moderate income': 'Moderate'
                , 'High income': 'High'
            }
            df_plot['Income Bracket'] = df_plot['Income Bracket'].map(var_map)
        if indicator == 'AFFH_3':
            var_map = {
                'Population 5 years and over who speak english "well" or "very well"': 'Speak english "well" or "very well"'
                , 'Population 5 years and over who speak english "not well" or "not at all"': 'Speak english "not well" or "not at all"'
            }
            df_plot['English Proficiency'] = df_plot['English Proficiency'].map(var_map)
        
        if 'Sort' in df_plot.columns:
            df_plot = df_plot.drop('Sort', axis=1)

    return df_plot



def make_fig(indicator, params, df_plot, county, jurisdiction):

    params = params['plot']

    if 'County' not in county:
        county = f'{county} County'

    color_discrete_map  = {
    
        'SACOG Region': '#9DC209'
        , county: '#1F45FC'
        , jurisdiction: '#FBB117'

        , 'American Indian or Alaska Native (NH)': '#E56717'
        , 'Native Hawaiian or other Pacific Islander (NH)': '#006A4E'
        , 'Other race or multiple races (NH)': '#7E587E'
        , 'Black or African American (NH)': '#FBB117'
        , 'Asian (NH)': '#9DC209'
        , 'Hispanic or Latino': '#1E90FF'
        , 'White (NH)': '#151B54'

        , 'American Indian or Alaska Native': '#E56717'
        , 'Native Hawaiian or other Pacific Islander': '#006A4E'
        , 'Other race or multiple races': '#7E587E'
        , 'Black or African American': '#FBB117'
        , 'Asian': '#9DC209'
        , 'Hispanic or Latino': '#1E90FF'
        , 'White (NH)': '#151B54'

        , 'American Indian or<br>Alaska Native': '#E56717'
        , 'Native Hawaiian or<br>other Pacific Islander': '#006A4E'
        , 'Other race or<br>multiple races': '#7E587E'
        , 'Black or<br>African American': '#FBB117'
        , 'Asian': '#9DC209'
        , 'Hispanic or<br>Latino': '#1E90FF'
        , 'White (NH)': '#151B54'

        , 'American Indian,<br>Alaska Native,<br>or Indigenous': '#E56717'
        , 'Asian or<br>Asian American': '#9DC209'
        , 'Black,<br>African American,<br>or African': '#FBB117'
        , 'Hispanic/Latina/e/o Only': '#1E90FF'
        , 'Native Hawaiian or<br>Other Pacific Islander': '#006A4E'
        , 'Other Race or<br>Multiple Races': '#7E587E'
        , 'White': '#151B54'

        , '2000': '#151B54'
        , '2010': '#1E90FF'
        , '2020': '#9DC209'
        , '2023': '#FBB117'
        , '2024': '#FBB117'

        , "2002":"#151B54"
        , "2007":"#1E90FF"
        , "2012":"#9DC209"
        , "2017":"#FBB117"
        , "2022":"#DC381F"

        , 'Same house': '#151B54'
        , 'Same city or town': '#1E90FF'
        , 'Same county': '#9DC209'
        , 'Elsewhere in CA': '#FBB117'
        , 'Elsewhere in U.S.': '#7E587E'
        , 'Abroad': '#DC381F'

        , 'Agriculture & Natural Resources': '#151B54'
        , 'Construction': '#1E90FF'
        , 'Manufacturing, Wholesale, & Transportation': '#9DC209'
        , 'Retail': '#FBB117'
        , 'Information': '#7E587E'
        , 'Finance & Professional Services': '#DC381F'
        , 'Health & Educational Services': '#008000'
        , 'Other': '#E56717'

        , 'Management, Business, Science, and Arts occupations': '#151B54'
        , 'Service occupations': '#1E90FF'
        , 'Sales and Office occupations': '#9DC209'
        , 'Natural Resources, Construction, and Maintenance occupations': '#FBB117'
        , 'Production, Transportation, and Material Moving occupations': '#7E587E'

        , 'Private company workers': '#151B54'
        , 'Self-employed workers': '#7E587E'
        , 'Private not-for-profit workers': '#1E90FF'
        , 'Local and state government workers': '#9DC209'
        , 'Federal government workers': '#FBB117'
        , 'Unpaid family workers': '#DC381F'

        , 'Place of residence': '#9DC209'
        , 'Place of work': '#1F45FC'

        , 'Agriculture & Natural Resources': '#151B54'
        , 'Arts, Recreation, & Other': '#E56717'
        , 'Construction': '#1E90FF'
        , 'Financial & Leasing': '#7FFFD4'
        , 'Government': "#906E3E"
        , 'Health & Educational Services': '#008000'
        , 'Information': '#7E587E'
        , 'Manufacturing & Wholesale': '#9DC209'
        , 'Professional & Managerial Services': '#CC7A8B'
        , 'Retail': '#FBB117'
        , 'Transportation & Utilities': '#620C4B'

        , 'Earnings &#36;1,250/month or less': '#151B54'
        , 'Earnings &#36;1,251/month to &#36;3,333/month': '#1E90FF'
        , 'Earnings greater than &#36;3,333/month': '#9DC209'

        , 'Renter occupied': '#9DC209'
        , 'Owner occupied': '#1F45FC'

        , 'Female-headed family': '#9DC209'
        , 'Male-headed family': '#1E90FF'
        , 'Married-couple family': '#151B54'
        , 'Other non-family': '#7E587E'
        , 'Single-person': '#FBB117'

        , 'No children': '#1F45FC'
        , 'One or more children under 18': '#9DC209'

        , "Occupied":"#1F45FC"
        , "Vacant": "#9DC209"

        , "For rent":"#151B54"
        , "For sale only":"#DC381F"
        , "For seasonal, recreational, or occasional use":"#1E90FF"
        , "Other vacant":"#9DC209"
        , "Rented, not occupied":"#7E587E"
        , "Sold, not occupied":"#FBB117"
        , "For migrant workers":"#006A4E"

        , "Less than &#36;250k":"#151B54"
        , "&#36;0k-&#36;250k":"#151B54"
        , "&#36;250k-&#36;500k":"#1E90FF"
        , "&#36;500k-&#36;750k":"#9DC209"
        , "&#36;750k-&#36;1M":"#FBB117"
        , "&#36;1M-&#36;1.5M":"#7E587E"
        , "&#36;1.5M-&#36;2M":"#DC381F"
        , "&#36;2M+":"#006A4E"

        , "Less than &#36;500":"#151B54"
        , "&#36;0-&#36;500":"#151B54"
        , "&#36;500-&#36;1,000":"#1E90FF"
        , "&#36;1,000-&#36;1,500":"#9DC209"
        , "&#36;1,500-&#36;2,000":"#FBB117"
        , "&#36;2,000-&#36;2,500":"#7E587E"
        , "&#36;2,500-&#36;3,000":"#DC381F"
        , "&#36;3,000 or more":"#006A4E"
        , "&#36;3,000+":"#006A4E"

        , 'Less than or equal to 1 person per room': '#E56717'
        , '1.01 to 1.5 occupants per room': '#9DC209'
        , 'More than 1.5 occupants per room': '#1F45FC'

        , 'Not computed': '#9B9A96'
        , '0%-30% of income used for housing': '#1F45FC'
        , '30%-50% of income used for housing': '#9DC209'
        , '50%+ of income used for housing': '#E56717'

        , "1 person households":"#151B54"
        , "2 person households":"#1E90FF"
        , "3-4 person households":"#9DC209"
        , "5 or more person households":"#FBB117"

        , 'Less than 30%': '#151B54'
        , '31%-50%': '#1E90FF'
        , '51%-80%': '#9DC209'
        , '81%-100%': '#FBB117'
        , 'Greater than 100%': '#DC381F'

        , 'Above poverty level': '#1F45FC'
        , 'Below poverty level': '#9DC209'

        , 'With a disability': '#9DC209'
        , 'No disability': '#1F45FC'

        , 'Employed': '#1F45FC'
        , 'Unemployed': '#9DC209'

        , "Sheltered - Emergency Shelter":"#151B54"
        , "Sheltered - Transitional Housing": "#1E90FF"
        , 'Unsheltered': "#9DC209"

        , "Emergency Shelter":"#151B54"
        , "Transitional Housing": "#1E90FF"
        , 'Unsheltered': "#9DC209"

        , 'Share of homeless population': '#9DC209'
        , 'Share of overall population': '#1F45FC'

        , 'Chronic Substance Abuse':"#151B54"
        , 'HIV/AIDS': "#1E90FF"
        , 'Severely Mentally Ill': "#9DC209"
        , 'Veterans': '#E56717'
        , 'Victims of Domestic Violence': '#FBB117'

        , "2020-21":"#151B54"
        , "2021-22":"#1E90FF"
        , "2022-23":"#9DC209"
        , "2023-24":"#FBB117"

        , '0%-30% of AMI': '#151B54'
        , '31%-50% of AMI': '#1E90FF'
        , '51%-80% of AMI': '#9DC209'
        , '81%-100% of AMI': '#FBB117'
        , 'Greater than 100% of AMI': '#DC381F'

        , 'Loan originated':'#1E90FF'
        , 'Application approved but not accepted':'#151B54'
        , 'Application denied': '#9DC209'
        , 'Application withdrawn by applicant':'#7FFFD4'
        , 'File closed for incompleteness':'#7E587E'
        , 'Purchased loan': '#FBB117'
        , 'Preapproval request denied':'#008000'
        , 'Preapproval request approved but not accepted':'#DC381F'

        , 'Speak english "well" or "very well"': '#1F45FC'
        , 'Speak english "not well" or "not at all"': '#9DC209'

    }

    df_plot = sort_plot(df_plot, indicator, jurisdiction)

    if indicator in ['HSG_6']:
        text_limit=0
    else:
        text_limit=5

    if params['text']:
        df_plot['text'] = np.where(df_plot[params['y']] >= text_limit, df_plot[params['y']].astype(str)+'%', '')


    # Line graphs
    if params['line']:
        fig = px.line(df_plot, x=params['x'], y=params['y'], color=params['color'], color_discrete_map=color_discrete_map, markers=True)
        if indicator == 'POPEMP_1':
            range_min = df_plot[params['y']].min()-4
            range_max = df_plot[params['y']].max()+4
            range_diff = abs(range_max-range_min)
            if range_diff <= 10: dtick=1
            elif (range_diff > 10) & (range_diff <= 50):   dtick=5
            elif (range_diff > 50) & (range_diff <= 100):  dtick=10
            elif (range_diff > 100) & (range_diff <= 200): dtick=25
            else: dtick=50
            fig.update_yaxes(ticksuffix='%', dtick=dtick, range=[range_min, range_max])
        if indicator in ['POPEMP_13', 'POPEMP_14']:
            if df_plot[params['y']].max() > 2.2:
                fig.update_yaxes(dtick=1, range=[0, 3.25])
            else:
                fig.update_yaxes(dtick=0.5, range=[0, 2.25])
        if indicator == 'POPEMP_15':
            fig.update_yaxes(ticksuffix='%', dtick=5, range=[0, 21])
        fig.update_xaxes(dtick=2, range=[df_plot[params['x']].min()-0.5, df_plot[params['x']].max()+0.5])

    # Bar graphs
    if params['bar']:
        if params['bar'] == 'stacked':
            if params['text']:
                fig = px.bar(df_plot, x=params['x'], y=params['y'], color=params['color'], color_discrete_map=color_discrete_map, text='text')
            else:
                fig = px.bar(df_plot, x=params['x'], y=params['y'], color=params['color'], color_discrete_map=color_discrete_map)
        if params['bar'] == 'group':
            if params['text']:
                fig = px.bar(df_plot, x=params['x'], y=params['y'], color=params['color'], color_discrete_map=color_discrete_map, barmode='group', text='text')
            else:
                fig = px.bar(df_plot, x=params['x'], y=params['y'], color=params['color'], color_discrete_map=color_discrete_map, barmode='group')
        if params['bar'] == 'simple':
            fig = px.bar(df_plot, x=params['x'], y=params['y'])
            fig.update_traces(marker_color='#1E90FF')

    if params['format'] == 'percent':
        if indicator in ['POPEMP_1', 'POPEMP_15', 'HSG_6', 'HSG_7', 'HSG_9', 'OVER_3', 'OVER_4', 'SEN_4', 'DISAB_1', 'HOMELS_2', 'ELI_3', 'ELI_4']:
            fig.update_yaxes(ticksuffix='%')
        else:
            fig.update_yaxes(dtick=25, ticksuffix='%', range=[0,102])
        fig.update_traces(textfont_color='white')
    if params['format'] == 'dollar':
        fig.update_yaxes(tickprefix='$', tickformat = ',.0f')
    if params['format'] == 'count':
        fig.update_yaxes(tickformat=',.0f')

    if indicator == 'POPEMP_10':
        fig.update_xaxes(tickvals=[0, 1, 2, 3, 4], ticktext=['Less than &#36;10k', '&#36;10k to &#36;25k', '&#36;25k to &#36;50k', '&#36;50k to &#36;75k', '&#36;75k or more'])
    if indicator == 'HOMELS_4':
        fig.update_layout(legend=dict(orientation="h", yanchor="bottom", y=-0.175,xanchor="right", x=0.75))

    if indicator not in ['POPEMP_4', 'POPEMP_10', 'HSG_1', 'LGFEM_5', 'HOMELS_3', 'HOMELS_4']:
        fig.update_layout(legend={'traceorder': 'reversed'})

    fig.update_yaxes(showgrid=True, gridwidth=1, gridcolor='lightgrey')
    fig.update_traces(hovertemplate="%{y}")

    return fig





def plot_rhna(fig, county, jurisdiction, indicator, params):

    PLOTLY_FONT_FAMILY='Microsoft YaHei'
    PLOTLY_TEMPLATE='plotly_white'
    PLOTLY_CONFIG={'modeBarButtonsToRemove': ['select', 'lasso', 'toImage'], 'displaylogo': False}

    path_plots = PATH_OUT / county.replace(' County', '') / jurisdiction / 'plots'

    fig.update_layout(
        legend_title=None
        , title=params['Title']
        , template=PLOTLY_TEMPLATE
        , font_family=PLOTLY_FONT_FAMILY
        , yaxis=dict(tickfont=dict(size=14))
        , xaxis=dict(tickfont=dict(size=14))
        )

    if indicator in ['POPEMP_3', 'POPEMP_5', 'POPEMP_6', 'POPEMP_7',
                        'POPEMP_8', 'POPEMP_9', 'POPEMP_16', 'POPEMP_23',
                        'POPEMP_24', 'HSG_2', 'HSG_3', 'HSG_7', 'HSG_9',
                        'OVER_2', 'OVER_7', 'LGFEM_2', 'DISAB_2', 'HOMELS_4',
                        'ELI_1', 'AFFH_3', 'FARM_2', 'LGFEM_1', 'SEN_2']:
        fig.update_layout(xaxis_title=None)
    
    if jurisdiction == 'Sacramento': fig.show(config=PLOTLY_CONFIG)
    
    if EXPORT:
        file_png  = path_plots / f'RHNA_{indicator}.png'
        try:
            fig.write_image(file=file_png , engine='kaleido', scale=1, width=1000, height=500)
        except: pass





def export_rhna_temp(county, jurisdiction, indicator):

    if EXPORT:

        FILE_YAML = load_yaml()

        path_juris = PATH_OUT / county.replace(' County', '') / jurisdiction / f'RHNA_{jurisdiction}.xlsx'

        if os.path.isfile(path_juris): wb = openpyxl.load_workbook(path_juris)
        else: wb = openpyxl.Workbook()

        if indicator not in wb.sheetnames:
            wb.create_sheet(indicator)
        if indicator in wb.sheetnames:
            sheet_index = wb.sheetnames.index(indicator)
            wb.remove(wb[indicator])
            wb.create_sheet(indicator, sheet_index)
        ws = wb[indicator]

        df_out = pd.DataFrame({'Temp': [FILE_YAML[indicator]['Title'], 'The data used for this indicator is not available for this jurisdiction']})
        if indicator == 'POPEMP_25': df_out = pd.DataFrame({'Temp': [FILE_YAML[indicator]['Title'], 'This indicator is still a work in progress']})

        for r in dataframe_to_rows(df_out, index=False, header=False): ws.append(r)
        ws['A1'].font = Font(bold=True, size=14)

        wb.save(path_juris)




def export_rhna(county, jurisdiction, indicator, params, df_prod, df_pct=None):

    if EXPORT:

        path_plots  = PATH_OUT / county.replace(' County', '') / jurisdiction / 'plots'
        path_tables = PATH_OUT / county.replace(' County', '') / jurisdiction / 'tables'
        path_plots .mkdir(parents=True, exist_ok=True)
        path_tables.mkdir(parents=True, exist_ok=True)

        try:

            try:
                path_plot = path_plots / f'RHNA_{indicator}.png'
                img = Image(path_plot)
            except:
                pass
            
            path_juris = PATH_OUT / county.replace(' County', '') / jurisdiction / f'RHNA_{jurisdiction}.xlsx'

            if os.path.isfile(path_juris): wb = openpyxl.load_workbook(path_juris)
            else: wb = openpyxl.Workbook()

            if indicator not in wb.sheetnames:
                wb.create_sheet(indicator)
            if indicator in wb.sheetnames:
                sheet_index = wb.sheetnames.index(indicator)
                wb.remove(wb[indicator])
                wb.create_sheet(indicator, sheet_index)
            ws = wb[indicator]

            # df_prod.columns = [col.title() for col in df_prod.columns if 'SACOG' not in col]
            # if df_pct is not None: df_pct.columns = [col.title() for col in df_pct.columns if 'SACOG' not in col]

            df_prod.to_csv(path_tables / f'{indicator}_prod.csv', index=False)
            if df_pct is not None: df_pct.to_csv(path_tables / f'{indicator}_pct.csv', index=False)

            df_prod2 = df_prod.reset_index().T.reset_index().T.drop(0, axis=1)
            if df_pct is not None:
                df_pct2 = df_pct.reset_index().T.reset_index().T.drop(0, axis=1)
            df_title = pd.DataFrame([['' for _ in range(len(df_prod2.columns))]])
            df_title.iloc[0, 0] = f'{indicator}: {params['Title']}'
            df_subtitle1 = pd.DataFrame([['' for _ in range(len(df_prod2.columns))]])
            if indicator in ['FARM_2']: df_subtitle1.iloc[0, 0] = f'Total: {county} County'
            elif indicator in ['HSG_11']: df_subtitle1.iloc[0, 0] = f'Total: {jurisdiction}'
            else: df_subtitle1.iloc[0, 0] = 'Total'
            if df_pct is not None:
                df_subtitle2 = pd.DataFrame([['' for _ in range(len(df_prod2.columns))]])
                df_subtitle2.iloc[0, 0] = 'Percent'

            df_space1 = pd.DataFrame([['' for _ in range(len(df_prod2.columns))]])
            df_space2 = pd.DataFrame([['' for _ in range(len(df_prod2.columns))]])

            df_source = pd.DataFrame([['' for _ in range(len(df_prod2.columns))]])
            df_source.iloc[0, 0] = 'Source'
            df_source.iloc[0, 1] = params['Source']

            df_years = pd.DataFrame([['' for _ in range(len(df_prod2.columns))]])
            df_years.iloc[0, 0] = 'Year(s)'
            df_years.iloc[0, 1] = params['Year(s)']

            df_notes = split_notes(indicator)

            df_title    .columns = df_prod2.columns
            df_subtitle1.columns = df_prod2.columns
            df_space1   .columns = df_prod2.columns
            df_space2   .columns = df_prod2.columns
            df_source   .columns = df_prod2.columns
            df_years    .columns = df_prod2.columns
            df_notes    .columns = df_prod2.columns[:2]
            if df_pct is not None: df_subtitle2.columns = df_prod2.columns

            if df_pct is not None: df_out = pd.concat([df_title, df_subtitle1, df_prod2, df_subtitle2, df_pct2, df_space1, df_space2, df_source, df_years, df_notes])
            else: df_out = pd.concat([df_title, df_subtitle1, df_prod2, df_space1, df_space2, df_source, df_years, df_notes])
            
            for r in dataframe_to_rows(df_out, index=False, header=False): ws.append(r)
                
            ws['A1'].font = Font(bold=True, size=14)
            ws['A2'].font = Font(bold=True, size=12)
            row_num = 3 + df_prod.shape[0] + 1
            ws[f'A{row_num}'].font = Font(bold=True, size=12) # 7 (style 1)

            for col in ws.columns:
                col = col[0].column_letter
                ws[f'{col}3'].border = Border(top=border_thin, left=border_thin, right=border_thin, bottom=border_thin)
                if df_pct is not None:
                    row_num = 3 + df_prod.shape[0] + 2
                    ws[f'{col}{row_num}'].border = Border(top=border_thin, left=border_thin, right=border_thin, bottom=border_thin) # 8 (style 1)

            format_numbers = '#,##0'
            format_percent = '0.0%'
            format_ratio   = '0.00'
            formet_dollars = '$#,##0'
            
            for col in ws.columns:
                max_length = 0
                column = col[0].column_letter
                if column == 'A':
                    pass
                else:
                    for cell in col:
                        try:
                            cell_num = re.search(r'.*?\.(.*)\>.*', str(cell))
                            cell_num = cell_num.group(1)
                            cell_num = cell_num[1:]
                            row_num = 3 + df_prod.shape[0]
                            if int(cell_num) > row_num or indicator in ['POPEMP_15'] or (indicator == 'POPEMP_1' and column in ['E', 'F', 'G']):# or (indicator in ['HSG_4'] and column == 'C'): # trying to get Percent column to show up properly
                                cell.value = float(cell.value)
                                cell.number_format = format_percent
                            else:
                                if indicator in ['POPEMP_13', 'POPEMP_14']:
                                    cell.number_format = format_ratio
                                else:
                                    cell.value = int(cell.value)
                                    cell.number_format = format_numbers
                        except: pass
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except: pass
                adjusted_width = (max_length + 1) * 1.1
                if column in ['A', 'B']:
                    if column == 'A': ws.column_dimensions[column].width = 20
                    else: ws.column_dimensions[column].width = 35
                else: ws.column_dimensions[column].width = adjusted_width

            try:
                row_num = df_out.shape[0] + 4
                ws.add_image(img, f'B{row_num}') # 21 (style 1)
            except: pass
            wb.save(path_juris)
        except Exception as e:
                    print(e); traceback.print_exc(); print()
                    export_rhna_temp(county, jurisdiction, indicator)















# def make_fig(indicator, df_plot, x, y, color, county, jurisdiction):

#     color_discrete_map  = {
    
#         'SACOG Region': '#9DC209'
#         , f'{county} County': '#1E90FF'
#         , jurisdiction: '#FBB117'

#         , 'American Indian or Alaska Native (NH)': '#E56717'
#         , 'Native Hawaiian or other Pacific Islander (NH)': '#006A4E'
#         , 'Other race or multiple races (NH)': '#7E587E'
#         , 'Black or African American (NH)': '#FBB117'
#         , 'Asian (NH)': '#9DC209'
#         , 'Hispanic or Latino': '#1E90FF'
#         , 'White (NH)': '#151B54'

#         , 'American Indian or<br>Alaska Native': '#E56717'
#         , 'Native Hawaiian or<br>other Pacific Islander': '#006A4E'
#         , 'Other race or<br>multiple races': '#7E587E'
#         , 'Black or<br>African American': '#FBB117'
#         , 'Asian': '#9DC209'
#         , 'Hispanic or<br>Latino': '#1E90FF'
#         , 'White (NH)': '#151B54'

#         , '2000': '#151B54'
#         , '2010': '#1E90FF'
#         , '2020': '#9DC209'
#         , '2023': '#FBB117'
#         , '2024': '#FBB117'

#         , 'Same house': '#151B54'
#         , 'Same city or town': '#1E90FF'
#         , 'Same county': '#9DC209'
#         , 'Elsewhere in CA': '#FBB117'
#         , 'Elsewhere in U.S.': '#7E587E'
#         , 'Abroad': '#DC381F'

#         , 'Agriculture & Natural Resources': '#151B54'
#         , 'Construction': '#1E90FF'
#         , 'Manufacturing, Wholesale, & Transportation': '#9DC209'
#         , 'Retail': '#FBB117'
#         , 'Information': '#7E587E'
#         , 'Finance & Professional Services': '#7FFFD4'
#         , 'Health & Educational Services': '#008000'
#         , 'Other': '#DC381F'

#         , 'Management, Business, Science, and Arts occupations': '#151B54'
#         , 'Service occupations': '#1E90FF'
#         , 'Sales and Office occupations': '#9DC209'
#         , 'Natural Resources, Construction, and Maintenance occupations': '#FBB117'
#         , 'Production, Transportation, and Material Moving occupations': '#7E587E'

#         , 'Private company workers': '#151B54'
#         , 'Self-employed workers': '#7E587E'
#         , 'Private not-for-profit workers': '#1E90FF'
#         , 'Local and state government workers': '#9DC209'
#         , 'Federal government workers': '#FBB117'
#         , 'Unpaid family workers': '#7FFFD4'

#         , 'Place of residence': '#9DC209'
#         , 'Place of work': '#151B54'

#         , 'Agriculture & Natural Resources': '#151B54'
#         , 'Arts, Recreation, & Other': '#E56717'
#         , 'Construction': '#1E90FF'
#         , 'Financial & Leasing': '#7FFFD4'
#         , 'Government': "#906E3E"
#         , 'Health & Educational Services': '#008000'
#         , 'Information': '#7E587E'
#         , 'Manufacturing & Wholesale': '#9DC209'
#         , 'Professional & Managerial Services': '#CC7A8B'
#         , 'Retail': '#FBB117'
#         , 'Transportation & Utilities': '#620C4B'

#         , 'Earnings &#36;1,250/month or less': '#151B54'
#         , 'Earnings &#36;1,251/month to &#36;3,333/month': '#1E90FF'
#         , 'Earnings greater than &#36;3,333/month': '#9DC209'

#         , 'Renter occupied': '#9DC209'
#         , 'Owner occupied': '#151B54'

#         , 'Female-headed family': '#9DC209'
#         , 'Male-headed family': '#1E90FF'
#         , 'Married-couple family': '#151B54'
#         , 'Other non-family': '#7E587E'
#         , 'Single-person': '#FBB117'

#         , 'No children': '#151B54'
#         , 'One or more children under 18': '#9DC209'

#         , "Occupied":"#151B54"
#         , "Vacant": "#9DC209"

#         , "For rent":"#151B54"
#         , "For sale only":"#DC381F"
#         , "For seasonal, recreational, or occasional use":"#1E90FF"
#         , "Other vacant":"#9DC209"
#         , "Rented, not occupied":"#7E587E"
#         , "Sold, not occupied":"#FBB117"
#         , "For migrant workers":"#006A4E"

#         , "Less than &#36;250k":"#151B54"
#         , "&#36;250k-&#36;500k":"#1E90FF"
#         , "&#36;500k-&#36;750k":"#9DC209"
#         , "&#36;750k-&#36;1M":"#FBB117"
#         , "&#36;1M-&#36;1.5M":"#7E587E"
#         , "&#36;1.5M-&#36;2M":"#DC381F"
#         , "&#36;2M+":"#006A4E"
#     }

#     if indicator in ['POPEMP_1', 'POPEMP_13', 'POPEMP_14', 'POPEMP_15']:
#         fig = px.line(df_plot, x=x, y=y, color=color, color_discrete_map=color_discrete_map, markers=True)
#         if indicator == 'POPEMP_1':
#             range_min = df_plot[y].min()-4
#             range_max = df_plot[y].max()+4
#             range_diff = abs(range_max-range_min)
#             if range_diff <= 10: dtick=1
#             elif (range_diff > 10) & (range_diff <= 50):   dtick=5
#             elif (range_diff > 50) & (range_diff <= 100):  dtick=10
#             elif (range_diff > 100) & (range_diff <= 200): dtick=25
#             else: dtick=50
#             fig.update_yaxes(ticksuffix='%', dtick=dtick, range=[range_min, range_max])
#         if indicator in ['POPEMP_13', 'POPEMP_14']:
#             if df_plot[y].max() > 2.2:
#                 fig.update_yaxes(dtick=1, range=[0, 3.25])
#             else:
#                 fig.update_yaxes(dtick=0.5, range=[0, 2.25])
#         if indicator == 'POPEMP_15':
#             fig.update_yaxes(ticksuffix='%', dtick=5, range=[0, 21])
#         fig.update_yaxes(showgrid=True, gridwidth=1, gridcolor='lightgrey')
#         fig.update_xaxes(dtick=2, range=[df_plot[x].min()-0.5, df_plot[x].max()+0.5])

#     elif indicator in ['POPEMP_4', 'POPEMP_10', 'POPEMP_11', 'POPEMP_12', 'POPEMP_20', 'POPEMP_21', 'HSG_1', 'HSG_5']:
#         if indicator in ['POPEMP_11', 'POPEMP_12']:
#             fig = px.bar(df_plot, x=x, y=y, color=color, color_discrete_map=color_discrete_map)
#         else:
#             fig = px.bar(df_plot, x=x, y=y, color=color, color_discrete_map=color_discrete_map, barmode='group')
#         if indicator in ['HSG_7']:
#             fig.update_yaxes(ticksuffix='%')
#         else:
#             fig.update_yaxes(tickformat=',.0f')
#         if indicator == 'POPEMP_10':
#             fig.update_xaxes(tickvals=[0, 1, 2, 3, 4], ticktext=['Less than &#36;10k', '&#36;10k to &#36;25k', '&#36;25k to &#36;50k', '&#36;50k to &#36;75k', '&#36;75k or more'])
#     elif indicator in ['HSG_4']:
#         fig = px.bar(df_plot, x=x, y=y)
#         fig.update_traces(marker_color='#1E90FF')
#         fig.update_yaxes(tickformat=',.0f')
#     else:
#         if indicator in ['POPEMP_23', 'POPEMP_24', 'HSG_2']:
#             if indicator == 'POPEMP_23':
#                 df_plot['Household Type'] = df_plot['Household Type'].str.replace(' households', '')
#             elif indicator == 'POPEMP_24':
#                 var_map = {
#                     'Households with no children': 'No children'
#                     , 'Households with 1 or more children under 18': "One or more children under 18"
#                 }
#                 df_plot['Household Type'] = df_plot['Household Type'].map(var_map)
#             elif indicator == 'HSG_2':
#                 var_map = {
#                     'Occupied housing units': 'Occupied',
#                     'Vacant housing units': 'Vacant'
#                 }
#                 df_plot['Occupancy Status'] = df_plot['Occupancy Status'].map(var_map)
#         if indicator in ['HSG_6']: text_limit=0
#         else: text_limit=5
#         df_plot['text'] = np.where(df_plot[y] >= text_limit, df_plot[y].astype(str)+'%', '')
#         fig = px.bar(df_plot, x=x, y=y, color=color, color_discrete_map=color_discrete_map, text='text')
#         if indicator in ['HSG_6']:
#             fig.update_yaxes(ticksuffix='%')
#         else:
#             fig.update_yaxes(dtick=25, ticksuffix='%', range=[0,102])
#         fig.update_traces(textfont_color='white')
    
#     fig.update_layout(legend={'traceorder': 'reversed'})
#     fig.update_yaxes(showgrid=True, gridwidth=1, gridcolor='lightgrey')

#     return fig  



